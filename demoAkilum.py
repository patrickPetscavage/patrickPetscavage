import os
import io
import uuid
import streamlit as st
from dotenv import load_dotenv
from openai import OpenAI

import chromadb
import pdfplumber
from docx import Document
import openpyxl

# Load environment variables
load_dotenv()

# Cached OpenAI client
@st.cache_resource
def get_client(api_key: str):
    return OpenAI(api_key=api_key)

# Persistent Chroma DB client
@st.cache_resource
def get_chroma():
    return chromadb.PersistentClient(path=".chroma_db")

chroma_client = get_chroma()

# UI reset helper (clears inputs)
def reset_ui():
    keys_to_clear = [
        "chat_input",
        "search_input",
        "temp_input",
    ]

    for key in keys_to_clear:
        if key in st.session_state:
            del st.session_state[key]

# Normalize company/project names
def normalize(name: str):
    return name.strip().lower().replace(" ", "_")


# Get or create company collection in Chroma
def get_company_collection(company_id: str):
    safe_name = f"company_{normalize(company_id)}"
    return chroma_client.get_or_create_collection(name=safe_name)

# File readers (PDF, DOCX, Excel, TXT)
def read_pdf(file):
    text = ""
    with pdfplumber.open(file) as pdf:
        for page in pdf.pages:
            text += (page.extract_text() or "") + "\n"
    return text

def read_docx(file):
    doc = Document(io.BytesIO(file.read()))
    return "\n".join([p.text for p in doc.paragraphs])

def read_excel(file):
    wb = openpyxl.load_workbook(io.BytesIO(file.read()), data_only=True)
    text = ""

    for sheet in wb.worksheets:
        text += f"\n## {sheet.title}\n"
        for row in sheet.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None]
            if vals:
                text += " | ".join(vals) + "\n"

    return text


# Read multiple uploaded files
def read_files(files):
    all_text = []

    for f in files:
        ext = f.name.split(".")[-1].lower()

        if ext == "pdf":
            txt = read_pdf(f)
        elif ext == "docx":
            txt = read_docx(f)
        elif ext in ["xlsx", "xlsm"]:
            txt = read_excel(f)
        else:
            txt = f.read().decode("utf-8", errors="ignore")

        all_text.append(f"[DOCUMENT: {f.name}]\n{txt}")

    return "\n\n".join(all_text)

# Split text into chunks for embedding
def split_text(text, size=400):
    words = text.split()
    return [" ".join(words[i:i+size]) for i in range(0, len(words), size)]


# Create embeddings using OpenAI
def embed(client, texts):
    res = client.embeddings.create(
        model="text-embedding-3-small",
        input=texts
    )
    return [r.embedding for r in res.data]

# Store chunks + embeddings in Chroma
def store_chunks(collection, project, chunks, client):
    embeddings = embed(client, chunks)
    ids = [str(uuid.uuid4()) for _ in chunks]

    collection.add(
        ids=ids,
        documents=chunks,
        embeddings=embeddings,
        metadatas=[{"project": project} for _ in chunks]
    )

# Retrieve relevant chunks from vector DB
def retrieve(collection, project, client, query, k=8):
    query_embedding = embed(client, [query])[0]

    results = collection.query(
        query_embeddings=[query_embedding],
        n_results=k,
        where={"project": project}
    )

    return results["documents"][0] if results["documents"] else []

# Generate final answer using retrieved context
def answer(client, question, context):
    prompt = f"""
You are a document assistant.

Use ONLY the context below.

CONTEXT:
{context}

QUESTION:
{question}
"""

    res = client.responses.create(
        model="gpt-4o-mini",
        input=prompt,
        max_output_tokens=300
    )

    return res.output[0].content[0].text

# Main Streamlit app
def main():
    st.title("🏢 Construction AI Bot")

    api_key = st.sidebar.text_input("OpenAI API Key", type="password")
    if not api_key:
        st.warning("Enter API key")
        return

    client = get_client(api_key)

    # COMPANY MANAGEMENT SECTION
    st.sidebar.header("🏢 Company")

    all_collections = chroma_client.list_collections()

    companies = [
        c.name.replace("company_", "")
        for c in all_collections
        if c.name.startswith("company_") and c.name != "company_"
    ]

    # FIX: ensure company state always exists BEFORE usage
    if "company" not in st.session_state:
        st.session_state["company"] = None

    new_company = st.sidebar.text_input("Create / Enter Company", key="company_input")

    if st.sidebar.button("➕ Set Company"):
        if new_company.strip():
            normalized = normalize(new_company)

            # IMPORTANT FIX: actually create collection immediately
            get_company_collection(normalized)

            st.session_state["company"] = normalized
            reset_ui()
            st.rerun()

    # FIX: keep dropdown stable and safe
    safe_companies = sorted(list(set(companies))) if companies else []

    selected_company = st.sidebar.selectbox(
        "Select Company",
        safe_companies if safe_companies else ["No companies yet"],
        key="company_select"
    )

    if selected_company != "No companies yet":
        normalized_selected = normalize(selected_company)

        if st.session_state.get("company") != normalized_selected:
            st.session_state["company"] = normalized_selected
            reset_ui()
            st.rerun()

    company = st.session_state["company"]

    # FIX: safety check so UI doesn't crash
    if company:
        st.sidebar.success(f"Active Company: {company}")
    else:
        st.sidebar.warning("No company selected")

    # FIX: prevent crash when no company exists
    if not company:
        st.stop()

    # FIX: collection must be created AFTER company is guaranteed
    collection = get_company_collection(company)

    # PROJECT MANAGEMENT SECTION
    def get_projects():
        data = collection.get(include=["metadatas"])
        metas = data.get("metadatas", [])

        return sorted(list(set(
            m["project"] for m in metas if m and "project" in m
        )))

    st.sidebar.header("📁 Project")

    new_project = st.sidebar.text_input("Create Project", key="project_input")

    if st.sidebar.button("➕ Set Project"):
        if new_project.strip():
            st.session_state["project"] = new_project
            reset_ui()
            st.rerun()

    projects = get_projects()

    selected_project = st.sidebar.selectbox(
        "Select Project",
        projects if projects else ["No projects yet"],
        key="project_select"
    )

    if selected_project != "No projects yet":
        if st.session_state.get("project") != selected_project:
            st.session_state["project"] = selected_project
            reset_ui()
            st.rerun()

    if "project" not in st.session_state:
        st.session_state["project"] = selected_project

    project = st.session_state["project"]
    st.sidebar.success(f"Active Project: {project}")

    # FILE UPLOAD + INDEXING SECTION
    files = st.sidebar.file_uploader(
        "Upload Documents",
        accept_multiple_files=True,
        key=f"uploader_{company}_{project}"
    )

    if st.sidebar.button("Index Documents"):
        if not files:
            st.error("Upload files first")
        else:
            text = read_files(files)
            chunks = split_text(text)

            store_chunks(collection, project, chunks, client)

            st.success(f"Indexed {len(chunks)} chunks")

    # CHAT MEMORY SETUP
    if "chat" not in st.session_state:
        st.session_state.chat = []

    for msg in st.session_state.chat:
        st.chat_message(msg["role"]).write(msg["text"])

    # CHAT INPUT + RAG FLOW
    prompt = st.chat_input("Ask something...", key="chat_input")

    if prompt:
        st.chat_message("user").write(prompt)

        context = retrieve(collection, project, client, prompt, k=8)
        context_text = "\n\n".join(context)

        reply = answer(client, prompt, context_text)

        st.chat_message("assistant").write(reply)

        st.session_state.chat.append({"role": "user", "text": prompt})
        st.session_state.chat.append({"role": "assistant", "text": reply})

#START APP
if __name__ == "__main__":
    main()

#Next upgrades/impovements
#Company database in sql
#Proejct database for each company in sql
#Better file system
#User authentication
